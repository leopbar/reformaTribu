"""Exportação dos itens aprovados (XLSX e CSV) com layout configurável, e relatório em PDF."""

from __future__ import annotations

import csv
import io
import uuid
from collections import Counter
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from jinja2 import Environment, FileSystemLoader, select_autoescape
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.analise import transicao
from app.core.codes import formatar_cnpj, formatar_codigo
from app.models import (
    Audit,
    AuditItem,
    CClassTribCode,
    Company,
    ExportJob,
    ExportLayout,
    ItemReview,
    Organization,
    RefSnapshot,
    RefVersion,
    User,
)
from app.models.enums import StatusRevisao

COLUNAS_DISPONIVEIS: dict[str, str] = {
    "codigo_interno": "Código interno",
    "descricao": "Descrição",
    "tipo": "Tipo",
    "ncm": "NCM",
    "nbs": "NBS",
    "ncm_anterior": "NCM anterior",
    "cst_ibs_cbs": "CST IBS/CBS",
    "cclasstrib": "cClassTrib",
    "tipo_tratamento": "Tratamento",
    "dispositivo_legal": "Dispositivo legal",
    "gtin": "GTIN",
    "cest": "CEST",
    "unidade": "Unidade",
    "status_auditoria": "Resultado da análise",
    "confianca": "Confiança",
    "hipotese": "Hipótese aplicada",
    "conclusao": "Fundamentação",
    "reducao_ibs": "Redução IBS (%)",
    "reducao_cbs": "Redução CBS (%)",
    "imposto_seletivo": "Imposto Seletivo",
    "vigencia": "Vigência",
    "cenario": "Cenário",
    "alterado": "Alterado?",
    "aprovado_por": "Aprovado por",
    "aprovado_em": "Aprovado em",
}

LAYOUT_PADRAO = [
    "codigo_interno",
    "descricao",
    "tipo",
    "ncm",
    "nbs",
    "cst_ibs_cbs",
    "cclasstrib",
    "dispositivo_legal",
    "ncm_anterior",
    "alterado",
    "aprovado_por",
    "aprovado_em",
]

ROTULO_STATUS = {
    "classificado": "Classificado",
    "aguardando_informacao": "Aguardando informação",
    "revisao_contador": "Revisão do contador",
    "revisao_especialista": "Revisão do especialista",
}
ROTULO_CENARIO = {"venda_consumidor": "Venda ao consumidor"}
ROTULO_IS = {
    "sujeito": "Sujeito",
    "na_origem": "Na origem (cobrado na fabricação/importação; não recolhe na revenda)",
    "nao_sujeito": "Não sujeito",
    "indefinido": "A confirmar",
}
TEMPLATES = Path(__file__).with_name("templates")


def _usuarios(session: Session, ids: set[uuid.UUID]) -> dict[uuid.UUID, str]:
    if not ids:
        return {}
    return {u.id: u.nome for u in session.scalars(select(User).where(User.id.in_(ids)))}


def _so_revende_seletivo(session: Session, audit: Audit | None) -> bool:
    """A empresa informou que não fabrica nem importa produtos do Imposto Seletivo (ADR 0029)."""
    from app.analise import fatos as fatos_mod
    from app.analise.avaliacao import FATO_IS

    empresa = session.get(Company, audit.company_id) if audit else None
    if empresa is None:
        return False
    fato = fatos_mod.fatos_empresa(session, empresa).get(FATO_IS)
    return fato is not None and fato.valor == "nao"


def linhas_aprovadas(session: Session, audit_id: uuid.UUID, ncm_formatado: bool) -> list[dict[str, Any]]:
    itens = list(
        session.scalars(
            select(AuditItem)
            .where(AuditItem.audit_id == audit_id, AuditItem.revisao_status == StatusRevisao.APROVADO)
            .order_by(AuditItem.linha)
        )
    )
    nomes = _usuarios(session, {i.revisado_por for i in itens if i.revisado_por})
    audit = session.get(Audit, audit_id)
    so_revende = _so_revende_seletivo(session, audit)
    vigencia = audit.data_referencia.strftime("%d/%m/%Y") if audit else ""
    saida = []
    for i in itens:
        eh_nbs = i.final_tipo_codigo == "nbs"
        codigo = (formatar_codigo(i.final_tipo_codigo, i.final_codigo) if ncm_formatado else i.final_codigo) or ""
        anterior = i.nbs if eh_nbs else i.ncm
        saida.append(
            {
                "codigo_interno": i.codigo_interno,
                "descricao": i.descricao,
                "tipo": "Serviço" if eh_nbs else "Produto",
                "ncm": "" if eh_nbs else codigo,
                "nbs": codigo if eh_nbs else "",
                "ncm_anterior": (formatar_codigo(i.final_tipo_codigo, anterior) if ncm_formatado else anterior) or "",
                "cst_ibs_cbs": i.final_cst or "",
                "cclasstrib": i.final_cclasstrib or "",
                "tipo_tratamento": (i.tipo_tratamento or "").replace("_", " "),
                "dispositivo_legal": i.final_dispositivo or "",
                "gtin": i.gtin or "",
                "cest": i.cest or "",
                "unidade": i.unidade or "",
                "status_auditoria": ROTULO_STATUS.get(i.status, i.status),
                "alterado": "Sim"
                if (
                    i.final_codigo != anterior
                    or (i.cst_atual and i.cst_atual != i.final_cst)
                    or (i.cclasstrib_atual and i.cclasstrib_atual != i.final_cclasstrib)
                )
                else "Não",
                "aprovado_por": "Automático (confiança alta)"
                if i.aprovado_automaticamente
                else (nomes.get(i.revisado_por, "") if i.revisado_por else ""),
                "aprovado_em": i.revisado_em.astimezone().strftime("%d/%m/%Y %H:%M") if i.revisado_em else "",
                "confianca": i.confianca_global or "",
                "hipotese": i.hipotese or "",
                "conclusao": i.conclusao or "",
                "reducao_ibs": float(i.perc_red_ibs) if i.perc_red_ibs is not None else "",
                "reducao_cbs": float(i.perc_red_cbs) if i.perc_red_cbs is not None else "",
                # A empresa que só revende não recolhe o IS (ADR 0029), mesmo em itens aprovados antes da resposta.
                "imposto_seletivo": ROTULO_IS.get(
                    "na_origem" if i.is_situacao == "sujeito" and so_revende else (i.is_situacao or ""), ""
                ),
                "vigencia": vigencia,
                "cenario": ROTULO_CENARIO.get(i.cenario, i.cenario),
                "descricao_normalizada": (i.identidade or {}).get("descricao_normalizada") or "",
                "descricao_oficial": (i.identidade or {}).get("descricao_oficial") or "",
                "identificacao": (i.identidade or {}).get("situacao") or "",
                "problemas_cadastro": ", ".join((i.identidade or {}).get("problemas_cadastro") or []),
            }
        )
    return saida


# Duas visões do resultado: o que o item é (cadastro enriquecido) e como ele é tributado (perfil).
COLUNAS_CADASTRO = [
    ("codigo_interno", "Código interno"),
    ("descricao", "Descrição original"),
    ("descricao_normalizada", "Descrição normalizada"),
    ("tipo", "Tipo"),
    ("ncm", "NCM"),
    ("nbs", "NBS"),
    ("descricao_oficial", "Descrição oficial do código"),
    ("ncm_anterior", "Código anterior"),
    ("identificacao", "Identificação"),
    ("problemas_cadastro", "Problemas no cadastro legado"),
    ("gtin", "GTIN"),
    ("unidade", "Unidade"),
]
COLUNAS_PERFIL = [
    ("codigo_interno", "Código interno"),
    ("descricao", "Descrição"),
    ("cenario", "Cenário"),
    ("vigencia", "Vigência"),
    ("cst_ibs_cbs", "CST IBS/CBS"),
    ("cclasstrib", "cClassTrib"),
    ("reducao_ibs", "Redução IBS (%)"),
    ("reducao_cbs", "Redução CBS (%)"),
    ("imposto_seletivo", "Imposto Seletivo"),
    ("hipotese", "Hipótese aplicada"),
    ("dispositivo_legal", "Fundamento"),
    ("conclusao", "Conclusão"),
    ("confianca", "Confiança"),
    ("aprovado_por", "Aprovado por"),
    ("aprovado_em", "Aprovado em"),
]


def _aba(ws: Any, linhas: list[dict[str, Any]], colunas: list[dict[str, str]]) -> None:
    ws.append([c["titulo"] for c in colunas])
    for cel in ws[1]:
        cel.font = Font(bold=True, color="FFFFFF")
        cel.fill = PatternFill("solid", fgColor="1F3A5F")
        cel.alignment = Alignment(vertical="center")
    for linha in linhas:
        ws.append([linha.get(c["campo"], "") for c in colunas])
    for idx, c in enumerate(colunas, start=1):
        largura = (
            min(
                60,
                max(len(c["titulo"]), *(len(str(linha.get(c["campo"], ""))) for linha in linhas[:500]))
                if linhas
                else len(c["titulo"]),
            )
            + 2
        )
        ws.column_dimensions[get_column_letter(idx)].width = largura
        # Códigos como texto para o Excel não remover zeros à esquerda.
        if c["campo"] in ("ncm", "nbs", "ncm_anterior", "cst_ibs_cbs", "cclasstrib", "gtin", "cest", "codigo_interno"):
            for cel in ws[get_column_letter(idx)][1:]:
                cel.number_format = "@"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def gerar_planilha(linhas: list[dict[str, Any]], colunas: list[dict[str, str]], formato: str, sep: str) -> bytes:
    if formato == "csv":
        texto = io.StringIO()
        w = csv.writer(texto, delimiter=sep, quoting=csv.QUOTE_MINIMAL, lineterminator="\r\n")
        w.writerow([c["titulo"] for c in colunas])
        for linha in linhas:
            w.writerow([linha.get(c["campo"], "") for c in colunas])
        return texto.getvalue().encode("utf-8-sig")
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Planilha ERP"
    _aba(ws, linhas, colunas)
    _aba(wb.create_sheet("Cadastro enriquecido"), linhas, [{"campo": c, "titulo": t} for c, t in COLUNAS_CADASTRO])
    _aba(wb.create_sheet("Perfil tributário"), linhas, [{"campo": c, "titulo": t} for c, t in COLUNAS_PERFIL])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def colunas_do_layout(layout: ExportLayout | None) -> tuple[list[dict[str, str]], str, bool]:
    if layout is None:
        return [{"campo": c, "titulo": COLUNAS_DISPONIVEIS[c]} for c in LAYOUT_PADRAO], ";", False
    cols = [c for c in layout.colunas if c.get("campo") in COLUNAS_DISPONIVEIS]
    return cols, layout.separador_csv or ";", layout.ncm_formatado


# ----------------------------------------------------------------------------- relatório --
def dados_relatorio(session: Session, audit_id: uuid.UUID) -> dict[str, Any]:
    audit = session.get(Audit, audit_id)
    assert audit is not None
    empresa = session.get(Company, audit.company_id)
    org = session.get(Organization, audit.org_id)
    snap = session.get(RefSnapshot, audit.snapshot_id) if audit.snapshot_id else None
    versoes = []
    if snap:
        for fonte, vid in [
            *(snap.versoes or {}).items(),
            *(("normas", v) for v in (snap.completude or {}).get("normas_versoes", [])),
        ]:
            v = session.get(RefVersion, uuid.UUID(vid)) if vid else None
            versoes.append(
                {
                    "fonte": fonte.upper(),
                    "rotulo": v.rotulo if v else "não importada",
                    "coletado_em": v.coletado_em.astimezone().strftime("%d/%m/%Y %H:%M") if v else "—",
                    "sha": v.sha256[:12] if v else "—",
                    "url": v.url_origem or "upload manual" if v else "—",
                }
            )
    itens = list(
        session.scalars(
            select(AuditItem)
            .where(AuditItem.audit_id == audit_id, AuditItem.ignorado.is_(False))
            .order_by(AuditItem.linha)
        )
    )
    status = Counter(i.status for i in itens)
    revisao = Counter(i.revisao_status for i in itens)
    motivos: Counter[str] = Counter(m for i in itens for m in i.motivos or [])
    alterados = [
        i
        for i in itens
        if i.revisao_status == StatusRevisao.APROVADO
        and (
            i.final_codigo != (i.nbs if i.final_tipo_codigo == "nbs" else i.ncm)
            or (i.cclasstrib_atual and i.cclasstrib_atual != i.final_cclasstrib)
        )
    ]
    revisoes = list(
        session.scalars(select(ItemReview).where(ItemReview.audit_id == audit_id).order_by(ItemReview.created_at))
    )
    por_revisor = Counter((r.user_email or "?", r.acao) for r in revisoes)
    cct_nomes: dict[str, str] = {}
    if snap and snap.versoes.get("cclasstrib"):
        for c in session.scalars(
            select(CClassTribCode).where(CClassTribCode.version_id == uuid.UUID(snap.versoes["cclasstrib"]))
        ):
            cct_nomes[c.codigo] = c.nome_reduzido or c.nome
    from app.pipeline.reasons import TEXTOS

    return {
        "org": org.nome if org else "",
        "empresa": empresa,
        "cnpj": formatar_cnpj(empresa.cnpj) if empresa else "",
        "audit": audit,
        "gerado_em": datetime.now(UTC).astimezone().strftime("%d/%m/%Y %H:%M"),
        "versoes": versoes,
        "regras_aprovadas": len(snap.regras_aprovadas) if snap else 0,
        "total": len(itens),
        "status": status,
        "revisao": revisao,
        "motivos": [(TEXTOS.get(m, (m, ""))[0], n) for m, n in motivos.most_common(12)],
        "alterados": [
            {
                "linha": i.linha,
                "codigo_interno": i.codigo_interno,
                "descricao": i.descricao,
                "antes": formatar_codigo(i.final_tipo_codigo, i.nbs if i.final_tipo_codigo == "nbs" else i.ncm) or "—",
                "depois": formatar_codigo(i.final_tipo_codigo, i.final_codigo),
                "cclasstrib": i.final_cclasstrib,
                "cct_nome": cct_nomes.get(i.final_cclasstrib or "", ""),
                "dispositivo": i.final_dispositivo or "",
            }
            for i in alterados[:2000]
        ],
        "total_alterados": len(alterados),
        "trilha": [{"revisor": k[0], "acao": k[1], "n": n} for k, n in sorted(por_revisor.items())],
        "custo": float(audit.custo_usd or 0),
        "aprovados_auto": sum(1 for i in itens if i.aprovado_automaticamente),
        "familias": len({i.thesis_id for i in itens if i.thesis_id}),
        "transicao": transicao.como_dict(audit.data_referencia),
        "tokens": audit.tokens or {},
        "configuracao": audit.configuracao or {},
    }


def gerar_pdf(session: Session, audit_id: uuid.UUID) -> bytes:
    from weasyprint import HTML

    env = Environment(loader=FileSystemLoader(TEMPLATES), autoescape=select_autoescape(["html"]))
    env.filters["num"] = lambda v: f"{int(v or 0):,}".replace(",", ".")
    html = env.get_template("relatorio.html").render(**dados_relatorio(session, audit_id))
    return HTML(string=html, base_url=str(TEMPLATES)).write_pdf()  # type: ignore[no-any-return]


def executar_job(session: Session, job: ExportJob) -> tuple[bytes, str]:
    audit = session.get(Audit, job.audit_id)
    assert audit is not None
    base = f"auditoria-{audit.nome[:40].strip().replace(' ', '_')}-{datetime.now().strftime('%Y%m%d-%H%M')}"
    if job.formato == "pdf":
        return gerar_pdf(session, job.audit_id), base + ".pdf"
    layout = session.get(ExportLayout, job.layout_id) if job.layout_id else None
    colunas, sep, ncm_fmt = colunas_do_layout(layout)
    linhas = linhas_aprovadas(session, job.audit_id, ncm_fmt)
    job.total_itens = len(linhas)
    return gerar_planilha(linhas, colunas, job.formato, sep), f"{base}.{job.formato}"
