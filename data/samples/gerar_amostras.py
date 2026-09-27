"""Gera as planilhas de exemplo (DADOS FICTÍCIOS) usadas em demonstrações e testes.

    uv run --project backend python data/samples/gerar_amostras.py

Os "erros propositais" (NCM sem zero à esquerda, sabonete líquido com NCM de sabonete em barra,
sucos sem informação de açúcar, GTIN inválido, duplicidades, códigos ausentes) servem para
demonstrar o sistema. NÃO use estes arquivos como conjunto-ouro: os códigos "corretos" aqui não
foram validados por um contador.
"""

from __future__ import annotations

import csv
import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

AQUI = Path(__file__).parent
AVISO = "DADOS FICTÍCIOS — gerados para demonstração e testes. Não correspondem a nenhuma empresa real."
rnd = random.Random(214)


def ean13(prefixo: str = "789", valido: bool = True) -> str:
    corpo = prefixo + "".join(str(rnd.randint(0, 9)) for _ in range(12 - len(prefixo)))
    soma = sum(int(c) * (3 if i % 2 else 1) for i, c in enumerate(corpo))
    dv = (10 - soma % 10) % 10
    if not valido:
        dv = (dv + 3) % 10
    return corpo + str(dv)


def salvar_xlsx(nome: str, cabecalho: list[str], linhas: list[list[object]], titulo: str) -> None:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Produtos"
    ws.append([titulo])
    ws.append([AVISO])
    ws.append([])
    ws.append(cabecalho)
    for c in ws[4]:
        c.font = Font(bold=True)
    for linha in linhas:
        ws.append(linha)
    wb.save(AQUI / nome)


# ------------------------------------------------------------------------- supermercado --
# (descrição abreviada como no ERP, NCM informado, categoria, unidade)
BASE_SUPER: list[tuple[str, str | int | None, str, str]] = [
    ("ARROZ TIPO 1 {m} 5KG", "10063021", "MERCEARIA", "PCT"),
    ("ARROZ PARBOILIZADO {m} 1KG", "10063011", "MERCEARIA", "PCT"),
    ("FEIJ CARIOCA {m} 1KG", "07133399", "MERCEARIA", "PCT"),
    ("FEIJ PRETO {m} 1KG", 7133319, "MERCEARIA", "PCT"),  # zero à esquerda perdido
    ("ACUC CRISTAL {m} 5KG", "17011400", "MERCEARIA", "PCT"),
    ("ACUC REFINADO {m} 1KG", "17019900", "MERCEARIA", "PCT"),
    ("CAF TORR MOIDO {m} 500G", "09012100", "MERCEARIA", "PCT"),
    ("LEITE UHT INTEG {m} 1L", 4012010, "LATICINIOS", "UN"),  # zero à esquerda perdido
    ("LEITE PO INTEG {m} 400G", "04022110", "LATICINIOS", "LAT"),
    ("MANT C/SAL {m} 200G", "04051000", "LATICINIOS", "UN"),
    ("QJO MUSSARELA FATIADO {m} KG", "04069030", "FRIOS", "KG"),
    ("IOG NATURAL {m} 170G", "04032000", "LATICINIOS", "UN"),
    ("OL SOJA {m} 900ML", "15079011", "MERCEARIA", "UN"),
    ("AZ EXTRA VIRGEM {m} 500ML", "15092000", "MERCEARIA", "UN"),
    ("FAR TRIGO TRAD {m} 1KG", "11010010", "MERCEARIA", "PCT"),
    ("MAC ESPAGUETE S/OVOS {m} 500G", "19021900", "MERCEARIA", "PCT"),
    ("MAC PARAFUSO C/OVOS {m} 500G", "19021100", "MERCEARIA", "PCT"),
    ("BISC RECH CHOC {m} 130G", "19021900", "BISCOITOS", "PCT"),  # NCM de massa: incoerente
    ("BISC AGUA E SAL {m} 200G", "19059020", "BISCOITOS", "PCT"),
    ("SAL REFINADO {m} 1KG", "25010020", "MERCEARIA", "PCT"),
    ("OVOS BRANCOS GRANDES {m} DZ", "04072100", "HORTIFRUTI", "DZ"),
    ("BANANA PRATA KG", "08039000", "HORTIFRUTI", "KG"),
    ("MACA GALA KG", "08081000", "HORTIFRUTI", "KG"),
    ("TOM ITALIANO KG", "07020000", "HORTIFRUTI", "KG"),
    ("BAT INGLESA KG", "07019000", "HORTIFRUTI", "KG"),
    ("CEB NACIONAL KG", "07031019", "HORTIFRUTI", "KG"),
    ("ALF CRESPA UN", "07051900", "HORTIFRUTI", "UN"),
    ("LARANJA PERA KG", "08051000", "HORTIFRUTI", "KG"),
    ("CENOURA KG", "07061000", "HORTIFRUTI", "KG"),
    ("CARNE BOV PATINHO RESF KG", "02013000", "ACOUGUE", "KG"),
    ("CARNE BOV MOIDA CONG {m} 500G", "02023000", "ACOUGUE", "PCT"),
    ("FGO INTEIRO CONG {m} KG", "02071210", "ACOUGUE", "KG"),
    ("LING TOSCANA {m} KG", "16010000", "ACOUGUE", "KG"),
    ("CAMARAO CINZA CONG {m} 400G", "03061790", "PEIXARIA", "PCT"),
    ("LAGOSTA CONG {m} KG", "03061200", "PEIXARIA", "KG"),  # exceção explícita do Anexo VII
    ("REFRIG COLA {m} 2L", "22021000", "BEBIDAS", "UN"),
    ("REFRIG GUARANA {m} 350ML LT", "22021000", "BEBIDAS", "LT"),
    ("AG MIN S/GAS {m} 500ML", "22011000", "BEBIDAS", "UN"),
    ("SUCO UVA {m} 1L", "20096100", "BEBIDAS", "UN"),  # não informa açúcar
    ("SUCO UVA INTEGRAL {m} 1,5L", "20096100", "BEBIDAS", "UN"),
    ("SUCO LARANJA C/ ACUC {m} 1L", "20091900", "BEBIDAS", "UN"),
    ("NECTAR PESSEGO {m} 1L", "20098990", "BEBIDAS", "UN"),  # néctar não é suco
    ("CERV PILSEN {m} 350ML LT", "22030000", "BEBIDAS", "LT"),
    ("VINH TINTO SECO {m} 750ML", "22042100", "BEBIDAS", "GAR"),
    ("CACHACA {m} 1L", "22084000", "BEBIDAS", "GAR"),
    ("CHOC AO LEITE {m} 90G", "18063210", "DOCES", "UN"),
    ("ACHOC PO {m} 400G", "18069000", "DOCES", "UN"),
    ("MARG C/SAL {m} 500G", "15171000", "LATICINIOS", "UN"),
    ("EXTR TOMATE {m} 340G", "20029000", "MERCEARIA", "UN"),
    ("MAION TRAD {m} 500G", "21039021", "MERCEARIA", "UN"),
    ("PAO FORMA TRAD {m} 500G", "19059010", "PADARIA", "PCT"),
    ("SAB BARRA ERVA DOCE {m} 90G", "34011190", "HIGIENE", "UN"),
    ("SAB LIQ ERVA DOCE {m} 250ML", "34011190", "HIGIENE", "UN"),  # líquido com NCM de barra
    ("CR DENT {m} 90G", "33061000", "HIGIENE", "UN"),
    ("ESC DENT MACIA {m}", "96032100", "HIGIENE", "UN"),
    ("PAP HIG FOLHA DUPLA {m} 12 ROLOS", "48181000", "HIGIENE", "PCT"),
    ("AG SANIT {m} 1L", "38089419", "LIMPEZA", "UN"),
    ("DET LIQ NEUTRO {m} 500ML", "34025000", "LIMPEZA", "UN"),
    ("SABAO EM PO {m} 1KG", "34025000", "LIMPEZA", "CX"),
    ("AMAC ROUPAS {m} 2L", "38099190", "LIMPEZA", "UN"),
    ("DESOD AEROSOL {m} 150ML", "33072010", "HIGIENE", "UN"),
    ("SHAMP {m} 350ML", "33051000", "HIGIENE", "UN"),
    ("FRALD INFANTIL M {m} 30UN", "96190000", "HIGIENE", "PCT"),
    ("CIGARRO {m} MACO", "24022000", "TABACARIA", "MC"),
    ("PROD DIVERSOS", None, "DIVERSOS", "UN"),  # descrição insuficiente
]
MARCAS = ["Vale Verde", "Boa Mesa", "Serra Azul", "Primor", "Campo Bom", "Estrela", "Solar", "Aurora"]


def supermercado() -> None:
    linhas: list[list[object]] = []
    cod = 1000
    for desc, ncm, cat, un in BASE_SUPER:
        variantes = 1 if "{m}" not in desc else rnd.randint(4, 6)
        for marca in rnd.sample(MARCAS, variantes) if variantes > 1 else [""]:
            cod += 1
            d = desc.replace("{m}", marca.upper()).replace("  ", " ").strip()
            gtin = ean13(valido=rnd.random() > 0.03) if cat != "HORTIFRUTI" else ""
            ncm_saida: object = ncm
            if ncm and rnd.random() < 0.04:
                ncm_saida = None  # alguns itens sem NCM cadastrado
            linhas.append([str(cod), d, ncm_saida, gtin, cat, un, marca or ""])
    # duplicidade proposital
    linhas.append(linhas[5][:])
    salvar_xlsx(
        "supermercado_ficticio.xlsx",
        ["Código", "Descrição do Produto", "NCM", "EAN", "Seção", "Unid.", "Marca"],
        linhas,
        "Supermercado Fictício Bom Preço — cadastro de produtos",
    )
    print(f"supermercado_ficticio.xlsx: {len(linhas)} itens")


# ------------------------------------------------------------------------------- padaria --
BASE_PADARIA: list[tuple[str, str | None]] = [
    ("PAO FRANCES KG", "19059090"),
    ("PAO DE SAL UN", None),
    ("PAO DE FORMA CASEIRO", "19059010"),
    ("PAO INTEGRAL FATIADO", "19059010"),
    ("PAO DE QJO UN", "19059090"),
    ("PAO DE QUEIJO CONG PCT 1KG", "19012090"),
    ("PAO DOCE C/ COCO", None),
    ("PAO DE LEITE", "19059090"),
    ("PAO SIRIO", "19059090"),
    ("BAGUETE", "19059090"),
    ("BROA DE MILHO", "19059090"),
    ("ROSCA DOCE", "19059090"),
    ("CROISSANT PRESUNTO E QJO", "19059090"),
    ("SONHO C/ CREME", None),
    ("BOLO CENOURA C/ COBERT CHOC", "19052090"),
    ("BOLO DE FUBA FATIA", None),
    ("BOLO ANIVERSARIO KG", "19059090"),
    ("TORTA DE FRANGO FATIA", "16023220"),
    ("COXINHA FRANGO UN", None),
    ("ESFIHA CARNE", None),
    ("PASTEL DE QUEIJO", None),
    ("SANDUICHE NATURAL", "16023220"),
    ("MISTO QUENTE", None),
    ("CAFE EXPRESSO XICARA", "09012100"),
    ("CAFE C/ LEITE COPO", None),
    ("CAPUCCINO", "21011200"),
    ("SUCO LARANJA NATURAL COPO 300ML", "20091200"),
    ("VITAMINA DE BANANA", None),
    ("LEITE INTEGRAL 1L", "04012010"),
    ("MANTEIGA POTE 200G", "04051000"),
    ("REQUEIJAO CREMOSO", "04061090"),
    ("PRESUNTO FATIADO KG", "16024100"),
    ("QUEIJO PRATO FATIADO KG", "04069020"),
    ("MORTADELA FATIADA KG", "16010000"),
    ("DOCE DE LEITE POTE", "19019020"),
    ("GELEIA MORANGO", "20079910"),
    ("BISCOITO AMANTEIGADO KG", "19053100"),
    ("BISCOITO DE POLVILHO PCT", "19059020"),
    ("TORRADA PCT", "19054000"),
    ("REFRIG LATA 350ML", "22021000"),
    ("AGUA MINERAL 500ML", "22011000"),
    ("OVOS DZ", "04072100"),
    ("SALGADO SORTIDO CENTO", None),
    ("QUICHE ALHO PORO", None),
    ("PUDIM FATIA", None),
    ("BRIGADEIRO UN", "17049020"),
    ("TRUFA CHOCOLATE", "18069000"),
    ("PANETONE 500G", "19052010"),
    ("CHOCOTONE 500G", "19052010"),
    ("CUCA DE BANANA", None),
]


def padaria() -> None:
    linhas: list[list[object]] = []
    for i, (desc, ncm) in enumerate(BASE_PADARIA, start=1):
        linhas.append([f"P{i:03d}", desc, ncm or "", "UN" if "KG" not in desc else "KG", "FABRICAÇÃO PRÓPRIA" if not ncm else "REVENDA"])
        # variações informais para chegar a ~100 itens
        if i % 2 == 0 or not ncm:
            informal = desc.lower().replace(" un", "").replace(" kg", "").capitalize() + " (balcão)"
            linhas.append([f"P{i:03d}B", informal, ncm or "", "UN", "FABRICAÇÃO PRÓPRIA"])
    with (AQUI / "padaria_ficticia.csv").open("w", encoding="cp1252", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow([AVISO])
        w.writerow(["COD", "PRODUTO", "CLASS FISCAL", "UN", "ORIGEM"])
        w.writerows(linhas)
    print(f"padaria_ficticia.csv: {len(linhas)} itens (Latin-1/Windows-1252, separador ';')")


# ------------------------------------------------------------------------------ farmácia --
BASE_FARMACIA: list[tuple[str, str, str]] = [
    ("DIPIRONA SODICA 500MG C/10 COMP", "30049099", "MEDICAMENTOS"),
    ("DIPIRONA GTS 20ML", "30049099", "MEDICAMENTOS"),
    ("PARACETAMOL 750MG C/20 COMP", "30049099", "MEDICAMENTOS"),
    ("IBUPROFENO 400MG C/10 CAPS", "30049099", "MEDICAMENTOS"),
    ("AMOXICILINA 500MG C/21 CAPS", "30041011", "MEDICAMENTOS"),
    ("AZITROMICINA 500MG C/3 COMP", "30042099", "MEDICAMENTOS"),
    ("LOSARTANA POTASSICA 50MG C/30", "30049099", "MEDICAMENTOS"),
    ("OMEPRAZOL 20MG C/28 CAPS", "30049099", "MEDICAMENTOS"),
    ("SINVASTATINA 20MG C/30", "30049099", "MEDICAMENTOS"),
    ("METFORMINA 850MG C/30", "30049099", "MEDICAMENTOS"),
    ("SORO FISIOLOGICO 0,9% 500ML", "30049099", "MEDICAMENTOS"),
    ("VITAMINA C 1G EFERV C/10", "21069030", "SUPLEMENTOS"),
    ("POLIVITAMINICO C/60 CAPS", "21069030", "SUPLEMENTOS"),
    ("XPE TOSSE INFANTIL 100ML", "30049099", "MEDICAMENTOS"),
    ("POM ASSADURA 45G", "30049099", "MEDICAMENTOS"),
    ("SERINGA 5ML C/AGULHA", "90183111", "DISPOSITIVOS"),
    ("AGULHA HIPODERMICA 25X7", "90183211", "DISPOSITIVOS"),
    ("LUVA LATEX PROCEDIMENTO CX100", "40151900", "DISPOSITIVOS"),
    ("TERMOMETRO DIGITAL", "90251990", "DISPOSITIVOS"),
    ("GAZE ESTERIL PCT 10UN", "30059090", "CURATIVOS"),
    ("ESPARADRAPO 10CMX4,5M", "30051090", "CURATIVOS"),
    ("CURATIVO ADESIVO C/40", "30051010", "CURATIVOS"),
    ("PRESERVATIVO C/3", "40141000", "DISPOSITIVOS"),
    ("TESTE GRAVIDEZ", "38221990", "DISPOSITIVOS"),
    ("MEDIDOR PRESSAO DIGITAL PULSO", "90189039", "DISPOSITIVOS"),
    ("FITA TESTE GLICEMIA C/50", "38221920", "DISPOSITIVOS"),
    ("CR DENT CLARIEADOR 70G", "33061000", "HIGIENE"),
    ("ESC DENT ULTRA MACIA", "96032100", "HIGIENE"),
    ("FIO DENTAL 50M", "33062000", "HIGIENE"),
    ("SAB LIQ ANTISSEPTICO 250ML", "34011190", "HIGIENE"),
    ("SAB BARRA GLICERINA 90G", "34011190", "HIGIENE"),
    ("FRALD GERIATRICA G C/8", "96190000", "HIGIENE"),
    ("ABS NOTURNO C/8", "96190000", "HIGIENE"),
    ("PAP HIG NEUTRO 4 ROLOS", "48181000", "HIGIENE"),
    ("PROT SOLAR FPS50 200ML", "33049990", "DERMOCOSMETICOS"),
    ("HIDRATANTE CORPORAL 400ML", "33049910", "DERMOCOSMETICOS"),
    ("PERFUME COLONIA 100ML", "33030020", "PERFUMARIA"),
    ("ESMALTE VERMELHO", "33043000", "PERFUMARIA"),
    ("SHAMPOO ANTICASPA 200ML", "33051000", "HIGIENE"),
    ("DESOD ROLL ON 50ML", "33072010", "HIGIENE"),
    ("REPEL LOCAO 100ML", "38089199", "HIGIENE"),
    ("ALCOOL GEL 70% 500ML", "22089000", "HIGIENE"),
    ("COTONETE C/75", "56012219", "HIGIENE"),
    ("LENCO UMEDECIDO C/50", "48189090", "HIGIENE"),
    ("MASC DESCARTAVEL C/50", "63079010", "DISPOSITIVOS"),
    ("ATADURA CREPOM 10CM", "30059020", "CURATIVOS"),
    ("BOLSA TERMICA GEL", "90191000", "DISPOSITIVOS"),
    ("INALADOR NEBULIZADOR", "90192010", "DISPOSITIVOS"),
    ("MULETA ALUMINIO PAR", "90211010", "DISPOSITIVOS"),
    ("MEIA COMPRESSAO 20-30", "61159500", "DISPOSITIVOS"),
]


def farmacia() -> None:
    linhas: list[list[object]] = []
    n = 0
    labs = ["LAB A", "LAB B", "GENERICO", "SIMILAR"]
    for desc, ncm, cat in BASE_FARMACIA:
        for lab in rnd.sample(labs, 3):
            n += 1
            linhas.append([n, f"{desc} {lab}", ncm, ean13("789", rnd.random() > 0.02), cat, ""])
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Cadastro"
    ws.append(["Id", "Descricao", "NCM", "Cod. Barras", "Grupo", "CST IBS/CBS atual"])
    for linha in linhas:
        ws.append(linha)
    ws2 = wb.create_sheet("Leia-me")
    ws2.append([AVISO])
    wb.save(AQUI / "farmacia_ficticia.xlsx")
    print(f"farmacia_ficticia.xlsx: {len(linhas)} itens")


# ------------------------------------------------------------------------------ serviços --
SERVICOS: list[tuple[str, str]] = [
    ("Consultoria em gestão empresarial — mensalidade", "1.1401.19.00"),
    ("Assessoria contábil mensal", "1.1302.21.00"),
    ("Auditoria de demonstrações financeiras", "1.1302.11.00"),
    ("Elaboração de declaração de imposto de renda PF", ""),
    ("Treinamento in company — gestão de estoques (8h)", ""),
    ("Curso online de Excel avançado", ""),
    ("Desenvolvimento de software sob encomenda (hora)", ""),
    ("Licenciamento de software de gestão — mensal", ""),
    ("Suporte técnico em informática (hora técnica)", ""),
    ("Hospedagem de site — plano anual", "1.1506.10.00"),
    ("Manutenção preventiva de ar-condicionado", ""),
    ("Instalação de rede elétrica predial", ""),
    ("Limpeza predial — contrato mensal", ""),
    ("Vigilância patrimonial — posto 12h", ""),
    ("Dedetização e controle de pragas", ""),
    ("Frete rodoviário de carga fracionada", "1.0501.12.10"),
    ("Transporte de passageiros por fretamento", ""),
    ("Locação de equipamentos de informática", ""),
    ("Serviços advocatícios — consultivo", "1.1301.20.00"),
    ("Tradução juramentada (lauda)", "1.1411.00.00"),
    ("Publicidade em redes sociais — gestão mensal", "1.1401.12.00"),
    ("Fotografia de produtos para e-commerce", ""),
    ("Consulta médica em clínica", ""),
    ("Exame laboratorial de sangue", ""),
    ("Sessão de fisioterapia", "1.2301.92.00"),
    ("Mensalidade escolar — ensino fundamental", "1.2201.20.00"),
    ("Curso de idiomas — mensalidade", ""),
    ("Serviço de estacionamento — diária", "1.0604.30.00"),
    ("Hospedagem em hotel — diária", ""),
    ("Serviço diversos", ""),
]


def servicos() -> None:
    linhas = [[f"S{i:03d}", d, n, "S"] for i, (d, n) in enumerate(SERVICOS, start=1)]
    salvar_xlsx(
        "servicos_ficticios.xlsx",
        ["Código", "Descrição do serviço", "NBS", "Tipo (P/S)"],
        linhas,
        "Consultoria Fictícia Exata — tabela de serviços",
    )
    print(f"servicos_ficticios.xlsx: {len(linhas)} itens")


if __name__ == "__main__":
    supermercado()
    padaria()
    farmacia()
    servicos()
